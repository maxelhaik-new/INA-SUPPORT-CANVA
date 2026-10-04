import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_scoping_document():
    doc = docx.Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.page_width = Inches(8.27)  # A4
        section.page_height = Inches(11.69)
        
    # Styling colors
    HEX_PRIMARY = "002B49"       # Bleu nuit INA
    HEX_ACCENT = "006699"        # Bleu acier
    HEX_DARK = "111827"          # Noir doux / anthracite
    HEX_MUTED = "4B5563"         # Gris soutenu
    HEX_LIGHT_BG = "F3F4F6"      # Gris clair sobre
    HEX_BORDER = "D1D5DB"        # Filet
    
    COLOR_PRIMARY = RGBColor(0, 43, 73)
    COLOR_ACCENT = RGBColor(0, 102, 153)
    COLOR_DARK = RGBColor(17, 24, 39)
    COLOR_MUTED = RGBColor(75, 85, 99)
    
    # Helper to set cell background and border
    def format_cell(cell, fill_hex=None, border_color=None, border_sz="4"):
        tcPr = cell._tc.get_or_add_tcPr()
        if fill_hex:
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
            tcPr.append(shd)
        if border_color:
            borders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:top w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
                    <w:left w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
                    <w:bottom w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
                    <w:right w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
                </w:tcBorders>
            ''')
            tcPr.append(borders)

    def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

    # ---------------- PAGE DE GARDE ----------------
    # Header tag
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(36)
    p_meta.paragraph_format.space_after = Pt(8)
    run_inst = p_meta.add_run("INSTITUT NATIONAL DE L'AUDIOVISUEL — DIRECTION DU DÉVELOPPEMENT & DE LA CRÉATION")
    run_inst.font.name = "Arial"
    run_inst.font.size = Pt(8.5)
    run_inst.font.bold = True
    run_inst.font.color.rgb = COLOR_ACCENT

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(6)
    run_title = p_title.add_run("NOTE DE CADRAGE STRATÉGIQUE")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(28)
    run_sub = p_sub.add_run("Rencontres Créatives de l'Audiovisuel et des Nouveaux Médias • Automne 2027")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_ACCENT

    # Divider bar table
    table_bar = doc.add_table(rows=1, cols=1)
    table_bar.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_bar.autofit = False
    table_bar.columns[0].width = Inches(6.47)
    format_cell(table_bar.rows[0].cells[0], fill_hex=HEX_PRIMARY)
    p_bar = table_bar.rows[0].cells[0].paragraphs[0]
    p_bar.paragraph_format.space_before = Pt(2)
    p_bar.paragraph_format.space_after = Pt(2)
    r_bar = p_bar.add_run("")
    r_bar.font.size = Pt(2)

    # Metadata card / Cartouche
    doc.add_paragraph().paragraph_format.space_after = Pt(36)
    card_table = doc.add_table(rows=4, cols=2)
    card_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    card_table.autofit = False
    card_table.columns[0].width = Inches(2.2)
    card_table.columns[1].width = Inches(4.27)

    meta_items = [
        ("Projet", "Rencontres Créatives de l'Audiovisuel et des Nouveaux Médias (Édition 2027)"),
        ("Émetteur", "Responsable Communication — Direction de la Communication & Création INA"),
        ("Destinataire", "Direction Générale & Direction des Partenariats Culturels"),
        ("Statut & Date", "Document de cadrage initial — Validation opérationnelle")
    ]

    for idx, (label, val) in enumerate(meta_items):
        cell_lbl = card_table.rows[idx].cells[0]
        cell_val = card_table.rows[idx].cells[1]
        
        format_cell(cell_lbl, fill_hex=HEX_LIGHT_BG, border_color=HEX_BORDER)
        format_cell(cell_val, fill_hex="FFFFFF", border_color=HEX_BORDER)
        set_cell_margins(cell_lbl, top=120, bottom=120, left=140, right=140)
        set_cell_margins(cell_val, top=120, bottom=120, left=140, right=140)

        p_l = cell_lbl.paragraphs[0]
        p_l.paragraph_format.space_before = Pt(2)
        p_l.paragraph_format.space_after = Pt(2)
        r_l = p_l.add_run(label)
        r_l.font.name = "Arial"
        r_l.font.size = Pt(9.5)
        r_l.font.bold = True
        r_l.font.color.rgb = COLOR_DARK

        p_v = cell_val.paragraphs[0]
        p_v.paragraph_format.space_before = Pt(2)
        p_v.paragraph_format.space_after = Pt(2)
        r_v = p_v.add_run(val)
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = COLOR_DARK

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(60)
    
    # Executive Quote block
    q_table = doc.add_table(rows=1, cols=1)
    q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    q_table.columns[0].width = Inches(6.47)
    cell_q = q_table.rows[0].cells[0]
    format_cell(cell_q, fill_hex="F9FAFB")
    tcPr = cell_q._tc.get_or_add_tcPr()
    borders_q = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_PRIMARY}"/>
            <w:top w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders_q)
    set_cell_margins(cell_q, top=140, bottom=140, left=180, right=180)
    
    pq = cell_q.paragraphs[0]
    pq.paragraph_format.space_before = Pt(4)
    pq.paragraph_format.space_after = Pt(4)
    rq = pq.add_run("« Fédérer 400 professionnels, créateurs indépendants et étudiants autour des mutations de l'image et du récit numérique, à travers un positionnement audacieux, contemporain et ouvert. »")
    rq.font.name = "Arial"
    rq.font.size = Pt(10)
    rq.font.italic = True
    rq.font.color.rgb = COLOR_MUTED

    # Page break after Cover page
    doc.add_page_break()

    # ---------------- CORPS DU DOCUMENT ----------------
    
    def add_section_header(num_str, title_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        
        r_num = p.add_run(f"{num_str}. ")
        r_num.font.name = "Arial"
        r_num.font.size = Pt(13)
        r_num.font.bold = True
        r_num.font.color.rgb = COLOR_ACCENT
        
        r_title = p.add_run(title_str.upper())
        r_title.font.name = "Arial"
        r_title.font.size = Pt(13)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_PRIMARY

        # Thin bottom border simulation using underline / small divider
        t_sep = doc.add_table(rows=1, cols=1)
        t_sep.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_sep.columns[0].width = Inches(6.47)
        format_cell(t_sep.rows[0].cells[0], fill_hex=HEX_BORDER)
        t_sep.rows[0].cells[0].paragraphs[0].paragraph_format.space_before = Pt(0)
        t_sep.rows[0].cells[0].paragraphs[0].paragraph_format.space_after = Pt(0)
        r_s = t_sep.rows[0].cells[0].paragraphs[0].add_run("")
        r_s.font.size = Pt(1)

    def add_body_p(text, bold_prefix=None, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK
        r_txt = p.add_run(text)
        r_txt.font.name = "Arial"
        r_txt.font.size = Pt(10)
        r_txt.font.color.rgb = COLOR_DARK
        return p

    def add_bullet_item(bold_label, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        r_b = p.add_run(bold_label + " : ")
        r_b.font.name = "Arial"
        r_b.font.size = Pt(10)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_DARK
        
        r_t = p.add_run(text)
        r_t.font.name = "Arial"
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = COLOR_DARK

    # SECTION 1 : CONTEXTE ET ENJEUX
    add_section_header("1", "Contexte & Enjeux Stratégiques")
    add_body_p(
        "Le paysage audiovisuel traverse une transformation sans précédent, marquée par l'émergence des technologies génératives, l'hybridation des formats immersifs et la fragmentation des usages chez les jeunes générations. Dans ce nouvel écosystème, l'INA — institution patrimoniale et pôle d'innovation des écritures — doit s'affirmer non seulement comme la mémoire vivante des images, mais comme l'accélérateur des créations de demain.",
        bold_prefix="Une mutation structurelle du récit : "
    )
    add_body_p(
        "Les « Rencontres Créatives de l'Audiovisuel et des Nouveaux Médias » programmées pour l'automne 2027 répondent à cette urgence. Ce rendez-vous inédit a pour vocation de décloisonner les disciplines, d'interroger les frontières techniques et artistiques, et de fonder un espace de dialogue prospectif entre pionniers du secteur, créateurs indépendants du web et talents émergents.",
        bold_prefix="L'opportunité pour l'INA : "
    )

    # SECTION 2 : 3 OBJECTIFS STRATÉGIQUES
    add_section_header("2", "Trois Objectifs Stratégiques Majeurs")
    add_body_p("Pour répondre avec rigueur à l'ambition fixée par la direction, le projet s'articule autour de trois objectifs directeurs mesurables :")
    
    add_bullet_item(
        "1. Positionner l'INA comme le carrefour prescripteur des nouveaux récits",
        "Réaffirmer le rôle de tiers de confiance et de laboratoire d'idées de l'INA en orchestrant la confrontation fertile entre audiovisuel traditionnel (cinéma, télévision, documentaire) et expressions numériques natives (formats verticaux, narrations immersives, IA générative appliquée à la production)."
    )
    add_bullet_item(
        "2. Fédérer et engager 400 acteurs clés de l'écosystème créatif",
        "Atteindre un visitorat qualifié et diversifié lors de l'édition d'automne 2027, en garantissant une parité d'expression entre décideurs de l'industrie, créateurs numériques émergents et vivier étudiant des filières créatives."
    )
    add_bullet_item(
        "3. Renouveler l'image de marque et dynamiser le réseau partenarial",
        "Déployer une prise de parole contemporaine, audacieuse et dénuée de tout jargon institutionnel austère, consolidant les partenariats institutionnels et ouvrant des opportunités de co-productions et de mécénat avec les industries créatives européennes."
    )

    # SECTION 3 : PUBLICS CIBLES
    add_section_header("3", "Cartographie des Publics Cibles (Jauge : 400 participants)")
    add_body_p("La réussite de l'événement repose sur un équilibre rigoureux entre trois collèges complémentaires :")

    # Table Publics
    table_pub = doc.add_table(rows=4, cols=3)
    table_pub.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_pub.autofit = False
    table_pub.columns[0].width = Inches(1.8)
    table_pub.columns[1].width = Inches(1.2)
    table_pub.columns[2].width = Inches(3.47)

    headers_pub = ["Segment de public", "Part / Jauge", "Attentes & Bénéfices recherchés"]
    for i, h in enumerate(headers_pub):
        c = table_pub.rows[0].cells[i]
        format_cell(c, fill_hex=HEX_PRIMARY, border_color=HEX_PRIMARY)
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    data_pub = [
        ("Professionnels de l'audiovisuel & décideurs", "45% (~180)", "Producteurs, diffuseurs, directeurs éditoriaux et distributeurs en quête de prospective, de nouveaux modèles économiques et d'analyses sectorielles solides."),
        ("Créateurs indépendants & talents web", "35% (~140)", "Vidéastes, scénaristes numériques, artistes 3D et narrateurs transmédias cherchant visibilité, partenariats financiers et légitimité institutionnelle."),
        ("Étudiants & jeunes diplômés d'écoles créatives", "20% (~80)", "Étudiants des formations audiovisuelles, écoles d'art, d'animation et du numérique (notamment INA sup) avides d'opportunités de mentorat et d'insertion professionnelle.")
    ]

    for row_idx, data in enumerate(data_pub, start=1):
        for col_idx, text in enumerate(data):
            c = table_pub.rows[row_idx].cells[col_idx]
            bg = HEX_LIGHT_BG if row_idx % 2 == 1 else "FFFFFF"
            format_cell(c, fill_hex=bg, border_color=HEX_BORDER)
            set_cell_margins(c, top=90, bottom=90, left=120, right=120)
            p = c.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
            r.font.color.rgb = COLOR_DARK

    # SECTION 4 : CONCEPT ÉDITORIAL & LIGNE DE COMMUNICATION
    add_section_header("4", "Concept Éditorial & Posture de Communication")
    add_body_p(
        "« Récits Vifs : L'Image Augmentée » (Titre de travail / Fil rouge). Le concept valorise la matière vivante des archives et de la création contemporaine qui s'hybrident sous l'impulsion du numérique. Loin d'une approche théorique ou nostalgique, l'angle choisi privilégie le « laboratoire » : des retours d'expériences concrets, des confrontations de méthodes et des démonstrations immersives.",
        bold_prefix="L'Angle éditorial : "
    )
    add_body_p(
        "Conformément aux directives de la direction, la communication rompt avec les codes institutionnels solennels. Elle privilégie un discours direct, visuel et percutant. L'audace éditoriale s'exprime par le refus de l'auto-célébration et par la mise en avant de questions vives : droit d'auteur à l'ère de l'IA, souveraineté culturelle des plateformes, nouveaux formats d'investigation.",
        bold_prefix="Tonalité & Ligne éditoriale : "
    )
    add_body_p(
        "Keynotes d'ouverture prospective, tables rondes pragmatiques (« Cas réels & Échecs constructifs »), masterclasses techniques et sessions de pitchs croisés étudiants-producteurs.",
        bold_prefix="Format des sessions : "
    )

    # SECTION 5 : FACTEURS CLÉS DE SUCCÈS
    add_section_header("5", "Facteurs Clés de Succès (FCS)")
    add_body_p("La réussite opérationnelle et l'impact d'image des Rencontres Créatives 2027 reposent sur 5 leviers critiques :")

    add_bullet_item("1. Rigueur du calendrier et anticipation des livrables", "Livraison dès ce vendredi de la note finale, de l'estimation budgétaire, du planning rétroactif, du deck partenaires (8 slides) et du kit visuel réseaux sociaux.")
    add_bullet_item("2. Mobilisation précoce des partenaires culturels", "Implication dès le S1 2027 des institutions tutélaires (CNC, Ministère de la Culture), des diffuseurs publics et des syndicats de producteurs pour co-construire la programmation.")
    add_bullet_item("3. Programmation inclusive et paritaire", "Garantie d'une représentativité équilibrée entre pionniers de la profession et figures de la création indépendante sur l'ensemble des panels.")
    add_bullet_item("4. Scénographie et expérience événementielle soignées", "Création d'un lieu d'échanges fluide favorisant le networking informel, doté d'une captation vidéo haut de gamme pour prolonger l'impact en digital.")
    add_bullet_item("5. Stratégie d'amplification sur les réseaux sociaux", "Plan de diffusion cross-plateforme (LinkedIn, Instagram, TikTok, YouTube) reposant sur des capsules vidéos dynamiques et des visuels graphiques audacieux.")

    # Save document
    output_path = "/Users/maximeelhaik/Documents/INA/INA SUPPORT CANVA/Note_de_cadrage_Rencontres_Creatives_INA_2027.docx"
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_scoping_document()
